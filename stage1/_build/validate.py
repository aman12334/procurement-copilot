"""Re-reads the generated vendor documents and checks them against the ground-truth files."""
import csv
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
GT = json.loads((ROOT / "ground_truth/canonical_vendor_bids.json").read_text())
QUAL = json.loads((ROOT / "ground_truth/vendor_qualification.json").read_text())
HIST = {r["line_id"]: float(r["fy2025_26_price_inr"]) for r in csv.DictReader(open(ROOT / "buyer/historical_reference_prices.csv"))}
RFX = {r["line_id"]: r for r in csv.DictReader(open(ROOT / "buyer/rfx_master.csv"))}
errors = []


def expect(vid):
    v = next(x for x in GT["vendors"] if x["vendor_id"] == vid)
    return {l["line_id"]: l for l in v["lines"] if l["quote_status"] != "not_quoted"}


def check(vid, found):
    exp = expect(vid)
    if set(found) != set(exp):
        errors.append(f"{vid}: coverage mismatch doc={sorted(set(found) ^ set(exp))}")
    for lid, val in found.items():
        e = exp.get(lid)
        if e and e["quoted_price"] is not None and abs(e["quoted_price"] - val) > 0.001:
            errors.append(f"{vid} {lid}: doc {val} != GT {e['quoted_price']}")
    print(f"{vid}: {len(found)} lines read from document, GT quoted {len(exp)}")


# V1 Excel
from openpyxl import load_workbook
ws = load_workbook(ROOT / "vendors/vendor_1/Deccan_Corrupack_Quotation_SCPL_RFQ014.xlsx")["Quotation"]
f = {}
for r in ws.iter_rows(min_row=10, values_only=True):
    if r[1] and str(r[1]).startswith("CP-"):
        f[r[1]] = float(re.match(r"[\d.]+", str(r[8])).group())
check("VND-1001", f)

# V2 PDF
from pypdf import PdfReader
txt = "\n".join(p.extract_text() for p in PdfReader(ROOT / "vendors/vendor_2/Shree_Ganesh_Packwell_Quotation_SGP-Q-2026-1187.pdf").pages)
f, dups = {}, {}
for m in re.finditer(r"(CP-\d{3})\n(?:[^\n]*\n){3}([\d,]+\.\d{2})\n(Nos|100 Nos|Set)\n(BLR|PNQ|HYD)", txt):
    lid, price = m.group(1), float(m.group(2).replace(",", ""))
    if lid in f:
        dups[lid] = price
    else:
        f[lid] = price
check("VND-1002", f)
assert "CP-011" in dups, "duplicate CP-011 missing"
txt = " ".join(txt.split())
for k in ("freight extra at actuals", "Special discount of 2%", "Payment:", "45 days", "Validity: 60 days"):
    assert k in txt, k

# V3 DOCX
from docx import Document
d = Document(ROOT / "vendors/vendor_3/Indus_Fibreboard_Offer_SCPL_RFQ014.docx")
f = {r.cells[0].text: float(r.cells[4].text) for r in d.tables[0].rows[1:]}
body = "\n".join(p.text for p in d.paragraphs)
for lid, usd in re.findall(r"(CP-0\d\d)[^.]*?USD ([\d.]+) per piece", body):
    f[lid] = float(usd)
check("VND-1003", f)
for k in ("90 days from date of offer",):
    assert any(k in c.text for t in d.tables for r in t.rows for c in r.cells), k
assert "valid until 31 December 2026" in body and "25,000 pieces" in body and "28 days" in body

# V4 HTML (source of the images)
html = (ROOT / "vendors/vendor_4/Sri_Vaigai_Rate_Card_Oct2026.html").read_text()
f = {}
for lid, rate in re.findall(r'data-line="(CP-\d{3})".*?<td class="r rate">(.*?)</td>', html):
    hand = re.search(r'class="hand">([\d.]+)<', rate)
    f[lid] = float(hand.group(1) if hand else re.match(r"[\d.]+", rate).group())
check("VND-1004", f)

# V5 email: per-kg rates and conversion
em = (ROOT / "vendors/vendor_5/Godavari_Packaging_email_response.txt").read_text()
for k in ("Rs 44/kg", "Rs 46/kg", "add Rs 3/kg", "Rs 5/kg extra", "Rs 41/kg", "same as last year's pricing", "Freight extra", "MOQ 5,000", "Payment 30 days"):
    assert k in em, k
for lid, e in expect("VND-1005").items():
    if e["quoted_uom"] == "per kg":
        calc = round(e["quoted_price"] * float(RFX[lid]["reference_weight_g"]) / 1000 + 1e-9, 2)
        if abs(calc - e["normalized_unit_price_inr"]) > 0.001:
            errors.append(f"V5 {lid} kg conversion")
    else:
        if e["normalized_unit_price_inr"] != HIST[lid]:
            errors.append(f"V5 {lid} last-year mismatch")
print(f"VND-1005: {len(expect('VND-1005'))} lines inferred from email")

# Normalization arithmetic for every record
for v in GT["vendors"]:
    for l in v["lines"]:
        if l["quote_status"] == "not_quoted" or l["quoted_price"] is None or l["quote_status"] == "quoted_ambiguous":
            continue
        n = round(l["quoted_price"] * l["fx_rate_to_inr"] / (l["pieces_per_quoted_unit"] or 1) + 1e-9, 2)
        if l["quoted_uom"] != "per kg" and abs(n - l["normalized_unit_price_inr"]) > 0.001:
            errors.append(f"norm {v['vendor_id']} {l['line_id']}")
        if abs(l["normalized_unit_price_inr"] + l["freight_inr_per_unit"] - l["landed_unit_price_inr"]) > 0.011:
            errors.append(f"landed {v['vendor_id']} {l['line_id']}")

# Qualification consistency
for q in QUAL["vendors"]:
    s = sum(a["points"] for a in q["answers"])
    ok = s >= 70 and not q["mandatory_failures"]
    if s != q["total_score"] or (q["status"] == "QUALIFIED") != ok:
        errors.append(f"qualification {q['vendor_id']}")
    print(f"{q['vendor_id']}: score {s} -> {q['status']}")

# No ground truth leakage into vendor docs
for p in (ROOT / "vendors").rglob("*"):
    if p.suffix in (".txt", ".html"):
        t = p.read_text()
        for word in ("normalized", "ground truth", "edge case", "HITL"):
            if word.lower() in t.lower():
                errors.append(f"leak '{word}' in {p.name}")

print("\nALL CHECKS PASSED" if not errors else "\n".join(["ERRORS:"] + errors))
