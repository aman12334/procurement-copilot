import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ground_truth import *
from edge_cases import build_manifest

ROOT = Path(__file__).resolve().parents[1]
BUY = ROOT / "buyer"
GT = ROOT / "ground_truth"
for d in (BUY, GT):
    d.mkdir(parents=True, exist_ok=True)


def dump(path, obj):
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


# ----------------------------------------------------------------- buyer
RFX_COLS = ["line_id", "item_description", "style", "dimensions_text", "dimension_basis", "ply", "flute", "board_construction",
            "combined_board_gsm", "min_bursting_factor_outer_liner", "min_bursting_strength_kg_cm2", "recycled_content_min_pct",
            "printing", "coating", "reference_weight_g", "delivery_plant", "annual_qty", "uom", "monthly_calloff_qty", "end_use"]
RFX_HEAD = ["Line ID", "Item Description", "Style", "Dimensions (mm)", "Dimension Basis", "Ply", "Flute", "Board Construction (GSM/BF per layer)",
            "Combined Board GSM (ref.)", "Min BF – Outer Liner", "Min Bursting Strength (kg/cm²)", "Min Recycled Content (%)",
            "Printing", "Coating / Special", "Reference Box Weight (g)", "Delivery Plant", "Annual Qty", "UoM", "Monthly Call-off (ref.)", "End Use"]


def write_rfx_master():
    with open(BUY / "rfx_master.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(RFX_COLS)
        for l in LINES:
            w.writerow([l[c] if l[c] is not None else "" for c in RFX_COLS])
    wb = Workbook()
    ws = wb.active
    ws.title = "RFx Lines"
    ws["A1"] = f"{BUYER['name']} – {EVENT['title']}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"RFx ID: {EVENT['rfx_id']}   |   Issued: {EVENT['issue_date']}   |   Response due: {EVENT['response_due_date']}   |   Quote basis: INR per piece/set, ex-GST, FOR delivery plant"
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="BFBFBF")
    for i, h in enumerate(RFX_HEAD, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r, l in enumerate(LINES, 5):
        for i, col in enumerate(RFX_COLS, 1):
            c = ws.cell(row=r, column=i, value=l[col])
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if col in ("annual_qty", "monthly_calloff_qty"):
                c.number_format = "#,##0"
    widths = [9, 46, 18, 16, 22, 5, 6, 48, 10, 9, 11, 10, 22, 34, 10, 9, 11, 6, 11, 34]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.row_dimensions[4].height = 45
    ws.freeze_panes = "C5"
    p = wb.create_sheet("Plants")
    p.append(["Code", "Name", "Address", "State"])
    for pl in PLANTS.values():
        p.append([pl["code"], pl["name"], pl["address"], pl["state"]])
    t = wb.create_sheet("Response Template")
    t.append([f["field"] for f in COMMERCIAL_FIELDS])
    for l in LINES:
        t.append([l["line_id"]] + [""] * (len(COMMERCIAL_FIELDS) - 1))
    wb.save(BUY / "rfx_master.xlsx")


BUYER_REQ = {
    "buyer": BUYER, "event": EVENT, "plants": list(PLANTS.values()),
    "scope": {"line_count": len(LINES), "lines_file": "buyer/rfx_master.csv",
              "annual_volume_total_units": sum(l["annual_qty"] for l in LINES),
              "annual_tonnage_reference_mt": round(sum(l["annual_qty"] * l["reference_weight_g"] for l in LINES) / 1e6, 1),
              "annual_volume_by_plant_units": {p: sum(l["annual_qty"] for l in LINES if l["delivery_plant"] == p) for p in PLANTS},
              "partial_bids": "Allowed. Vendors may quote a subset of lines; unquoted lines must be marked 'Not quoted'."},
    "requirements": [
        {"id": "BR-01", "area": "Delivery", "requirement": "Deliver to the plant specified per line (BLR / PNQ / HYD) as per weekly call-off schedule.", "type": "mandatory"},
        {"id": "BR-02", "area": "Annual volume", "requirement": "Annual quantities in rfx_master are estimates (+/-20%); no take-or-pay commitment.", "type": "information"},
        {"id": "BR-03", "area": "Lead time", "requirement": "Lead time from PO <= 14 days for repeat items; <= 21 days for first supply of new artwork / tooling.", "type": "preferred"},
        {"id": "BR-04", "area": "Quality – board", "requirement": "Meet or exceed the per-line board construction, BF, bursting strength and recycled-content minimums. Moisture content 7-9% at dispatch.", "type": "mandatory"},
        {"id": "BR-05", "area": "Quality – inspection", "requirement": "COA with every lot (GSM, BS, ECT/BCT, moisture; Cobb for coated lines). SCPL incoming AQL 2.5 (General II) on dimensions and print.", "type": "mandatory"},
        {"id": "BR-06", "area": "Quality – moisture resistance", "requirement": "Lines CP-012 and CP-029 require water-resistant coating meeting the Cobb value in rfx_master. Uncoated offers are non-compliant.", "type": "mandatory"},
        {"id": "BR-07", "area": "Certification", "requirement": "ISO 9001:2015 certification of supplying plant(s).", "type": "mandatory"},
        {"id": "BR-08", "area": "Certification", "requirement": "FSC Chain-of-Custody certification.", "type": "preferred"},
        {"id": "BR-09", "area": "Certification", "requirement": "BRCGS Packaging / ISO 22000 for lines used in food categories.", "type": "preferred"},
        {"id": "BR-10", "area": "Freight", "requirement": "Prices FOR SCPL delivery plant (freight included). Where freight is excluded or 'at actuals', SCPL applies a 5% evaluation loading on the unit price for comparison.", "type": "preferred", "evaluation_rule": {"freight_loading_pct": EVENT["freight_loading_pct"]}},
        {"id": "BR-11", "area": "Taxes", "requirement": "Quote exclusive of GST; state GST rate and HSN separately. Evaluation is ex-GST (input credit available).", "type": "mandatory"},
        {"id": "BR-12", "area": "Currency", "requirement": "Quote in INR. Non-INR quotes are accepted but converted at the fixed evaluation FX rate (1 USD = INR 88.50) for comparison; FX risk sits with the vendor in the contract.", "type": "preferred"},
        {"id": "BR-13", "area": "Payment terms", "requirement": "SCPL standard: 60 days from GRN. Minimum acceptable 45 days; shorter terms are a commercial deviation.", "type": "preferred", "standard_days": 60, "minimum_days": 45},
        {"id": "BR-14", "area": "Price validity", "requirement": "Prices firm for 180 days from quotation date. Minimum acceptable 90 days; shorter validity is a commercial deviation.", "type": "preferred", "standard_days": 180, "minimum_days": 90},
        {"id": "BR-15", "area": "MOQ", "requirement": "MOQ per line should not exceed one month's call-off quantity (annual / 12).", "type": "preferred"},
        {"id": "BR-16", "area": "Tooling", "requirement": "State one-time printing plate / cutting-die costs explicitly; tooling becomes SCPL property.", "type": "mandatory"},
        {"id": "BR-17", "area": "Packaging quality", "requirement": "Bundles of 25/50 strapped, palletised and stretch-wrapped; no warp > 5 mm; print registration +/- 1 mm; Pantone match per approved shade card.", "type": "mandatory"},
        {"id": "BR-18", "area": "Service level", "requirement": "OTIF >= 95% measured monthly; complaint acknowledgement within 24 h and CAPA within 7 days.", "type": "preferred"},
        {"id": "BR-19", "area": "Technical qualification", "requirement": "Score >= 70/100 on the technical questionnaire AND full compliance on all mandatory questions.", "type": "mandatory"},
        {"id": "BR-20", "area": "Price escalation", "requirement": "Any kraft-paper price variation clause must reference a published index (e.g., monthly kraft paper index) with a defined trigger and formula.", "type": "preferred"},
    ],
    "evaluation": {
        "comparison_basis": EVENT["normalization_basis"],
        "landed_price_formula": "normalized_unit_price_inr + freight (quoted, or 5% loading where excluded)",
        "unit_conversions": {"per piece": 1, "per box (carton itself)": 1, "per set (partitions)": 1, "per 100 nos": "divide by 100",
                              "per kg": "multiply by reference_weight_g / 1000 from rfx_master"},
        "weights": {"technical": 40, "commercial": 60},
        "award_strategy": "Line-wise or plant-wise award possible; maximum 3 vendors overall.",
    },
}

COMMERCIAL_FIELDS = [
    {"field": "line_id", "description": "SCPL line ID from rfx_master", "required": True},
    {"field": "unit_price", "description": "Basic price per quoted UoM, ex-GST", "required": True},
    {"field": "quoted_uom", "description": "Preferred: per piece (per set for partitions). Other UoMs must be clearly stated.", "required": True},
    {"field": "currency", "description": "INR preferred", "required": True},
    {"field": "moq", "description": "Minimum order quantity per release, in pieces", "required": True},
    {"field": "freight", "description": "Included (FOR plant) / amount per piece / excluded", "required": True},
    {"field": "taxes", "description": "GST rate and HSN code (prices must be exclusive of GST)", "required": True},
    {"field": "lead_time_days", "description": "Days from PO to delivery at plant", "required": True},
    {"field": "payment_terms", "description": "Credit days and trigger (GRN / invoice)", "required": True},
    {"field": "price_validity", "description": "Days or date until which price is firm", "required": True},
    {"field": "tooling_cost", "description": "One-time plates / dies cost, if any", "required": False},
    {"field": "annual_rebate", "description": "Volume rebate %, threshold and settlement", "required": False},
    {"field": "delivery_location", "description": "SCPL plant code (BLR / PNQ / HYD)", "required": True},
    {"field": "remarks", "description": "Deviations from specification or terms", "required": False},
]


def write_buyer():
    write_rfx_master()
    dump(BUY / "buyer_requirements.json", BUYER_REQ)
    dump(BUY / "technical_questionnaire.json", {"rfx_id": EVENT["rfx_id"], "questions": QUESTIONS, "qualification": QUALIFICATION, "mandatory_qids": MANDATORY_QIDS})
    dump(BUY / "commercial_terms.json", {
        "rfx_id": EVENT["rfx_id"], "response_fields": COMMERCIAL_FIELDS,
        "buyer_standard_terms": {"incoterm_equivalent": "FOR SCPL plant (delivered)", "currency": "INR", "payment": "60 days from GRN",
                                 "price_validity_days": 180, "taxes": "Exclusive of GST", "lead_time_days": 14,
                                 "moq_rule": "<= monthly call-off", "evaluation_fx_usd_inr": EVENT["evaluation_fx"]["USD_INR"],
                                 "freight_loading_pct_if_excluded": EVENT["freight_loading_pct"]},
        "deviation_rules": [
            {"term": "payment_days", "acceptable_min": 45, "below_min": "commercial deviation, flagged"},
            {"term": "price_validity_days", "acceptable_min": 90, "below_min": "commercial deviation, flagged"},
            {"term": "freight", "if_excluded": "apply 5% loading for comparison"},
            {"term": "conditional discounts/rebates", "treatment": "excluded from normalized price; shown as scenario"},
            {"term": "tooling", "treatment": "one-time cost reported separately; not amortised into unit price"},
        ]})
    with open(BUY / "historical_reference_prices.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["line_id", "item_description", "delivery_plant", "uom", "fy2025_26_price_inr", "price_basis", "fy2025_26_volume", "last_year_supplier", "implied_rate_inr_per_kg", "reference_weight_g"])
        for l in LINES:
            w.writerow([l["line_id"], l["item_description"], l["delivery_plant"], l["uom"], f"{l['historical_price_inr']:.2f}",
                        "Weighted avg PO price, delivered (FOR plant), ex-GST", round(l["annual_qty"] * 0.93),
                        LAST_YEAR_SUPPLIER[l["line_id"]], HIST_RATE_PER_KG[l["price_category"]], l["reference_weight_g"]])


# ----------------------------------------------------------------- ground truth
def deviations(vid):
    t = TERMS[vid]
    dev = []
    if t["payment_days"] < 45:
        dev.append(f"Payment {t['payment_days']} days < 45-day minimum")
    elif t["payment_days"] < 60:
        dev.append(f"Payment {t['payment_days']} days below 60-day standard (acceptable)")
    if t["price_validity_days"] < 90:
        dev.append(f"Price validity {t['price_validity_days']} days < 90-day minimum")
    elif t["price_validity_days"] < 180:
        dev.append(f"Price validity {t['price_validity_days']} days below 180-day standard (acceptable)")
    if t["freight_treatment"] in ("excluded", "included_PNQ_only"):
        dev.append("Freight not fully included – 5% evaluation loading applied")
    return dev


VENDOR_HITL = {
    "VND-1001": ["CP-026 possible price outlier"],
    "VND-1002": ["Duplicate CP-011 with conflicting prices", "Conditional 2% discount buried in note (scenario only)", "Freight 'at actuals' for BLR/HYD buried in footnote", "Tooling at actuals – amount unknown"],
    "VND-1003": ["Price validity conflict: table 90 days vs note 31-Dec-2026", "CP-014 description/dimension mismatch", "USD pricing on export lines (FX exposure)"],
    "VND-1004": ["Ambiguous handwritten price CP-010", "Fails technical qualification (no ISO 9001, Q05 partial, no BCP)", "CP-012 quoted without mandatory WR coating", "Image-based extraction – confidence review", "Payment 15 days and validity 30 days below minimums"],
    "VND-1005": ["Price by reference to last year's pricing (CP-016, CP-017, CP-025)", "Unclear small-run surcharge and escalation clause", "Incomplete coverage (20/30)", "Per-kg pricing relies on reference box weight", "Payment 30 days below 45-day minimum"],
}


def write_ground_truth():
    summ = vendor_summary()
    bids = {"rfx_id": EVENT["rfx_id"], "normalization_basis": EVENT["normalization_basis"], "evaluation_fx_usd_inr": FX,
            "freight_loading_pct": EVENT["freight_loading_pct"], "vendors": []}
    qual_out = {"rfx_id": EVENT["rfx_id"], "qualification_rule": QUALIFICATION, "mandatory_qids": MANDATORY_QIDS, "vendors": []}
    for v in VENDORS:
        vid = v["vendor_id"]
        qual = qualification(vid)
        recs = [r for r in RECORDS if r["vendor_id"] == vid]
        bids["vendors"].append({
            "vendor_id": vid, "vendor_name": v["vendor_name"], "source_document_type": SOURCE[vid][0], "source_file": SOURCE[vid][1],
            "header_terms": TERMS[vid], "commercial_deviations": deviations(vid),
            "technical_score": qual["total_score"], "technical_qualification_status": qual["status"],
            "mandatory_failures": qual["mandatory_failures"],
            "summary": summ[vid], "vendor_level_hitl_required": True, "vendor_level_hitl_reasons": VENDOR_HITL[vid],
            "questionnaire_answers": {a["qid"]: {"answer": a["vendor_answer"], "level": a["answer_level"], "points": a["points"]} for a in qual["answers"]},
            "lines": recs,
        })
        qual_out["vendors"].append({"vendor_id": vid, "vendor_name": v["vendor_name"], **qual,
                                    "missing_mandatory_certification": "ISO 9001:2015" if vid == "VND-1004" else None,
                                    "commercial_deviations": deviations(vid)})
    all5 = [v["vendor_id"] for v in VENDORS]
    qualified = [v["vendor_id"] for v in VENDORS if qualification(v["vendor_id"])["status"] == "QUALIFIED"]
    bids["evaluation_answers"] = {
        "technically_strongest": max(all5, key=lambda x: qualification(x)["total_score"]),
        "failed_qualification": [x for x in all5 if x not in qualified],
        "qualified_vendors": qualified,
        "common_basket_all_vendors": basket(all5),
        "common_basket_qualified_vendors": basket(qualified),
        "common_basket_full_coverage_qualified_vendors": basket(["VND-1001", "VND-1002", "VND-1003"]),
        "cheapest_overall_ignoring_qualification": basket(all5)["cheapest"],
        "cheapest_qualified": basket(qualified)["cheapest"],
        "line_wise_lowest_landed_qualified": line_winners(qualified),
        "line_wise_lowest_landed_all": line_winners(all5),
        "notes": [
            "V4 (Sri Vaigai) is cheapest on every basket but fails technical qualification – 'strong price, poor technical' case.",
            "V2 (Shree Ganesh) is cheapest among qualified vendors; if awarded > INR 1.5 Cr annual value its conditional 2% discount would further reduce price (not applied in normalized values).",
            "V1 totals are inflated by the CP-026 outlier (Rs %.2f vs hist Rs %.2f)." % (QUOTES['VND-1001']['CP-026']['price'], LINE_BY_ID['CP-026']['historical_price_inr']),
            "V3 (Indus) is technically strongest (94/100) but highest priced.",
        ],
    }
    dump(GT / "canonical_vendor_bids.json", bids)
    dump(GT / "vendor_qualification.json", qual_out)
    cols = ["vendor_id", "vendor_name", "line_id", "rfx_description", "delivery_plant", "annual_qty", "quote_status", "source_document_type",
            "source_location", "raw_price_text", "quoted_price", "quoted_currency", "quoted_uom", "pieces_per_quoted_unit", "fx_rate_to_inr",
            "normalized_unit_price_inr", "freight_treatment", "freight_inr_per_unit", "freight_is_estimated", "landed_unit_price_inr",
            "historical_price_inr", "variance_vs_historical_pct", "moq", "moq_exceeds_monthly_calloff", "lead_time_days", "payment_terms",
            "price_validity_until", "tax_treatment", "extraction_confidence", "comparable", "hitl_required", "hitl_reasons", "ambiguity", "edge_case_ids"]
    with open(GT / "normalized_line_items.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for r in RECORDS:
            row = []
            for c in cols:
                val = r.get(c)
                if isinstance(val, list):
                    val = "; ".join(val)
                row.append("" if val is None else val)
            w.writerow(row)
    dump(GT / "edge_case_manifest.json", build_manifest())


if __name__ == "__main__":
    write_buyer()
    write_ground_truth()
    print("buyer + ground truth written")
