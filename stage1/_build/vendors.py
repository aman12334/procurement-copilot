"""Writes the five messy vendor-facing response documents. No ground truth is written here."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ground_truth import V1_FIRST_ROW, V1_ROW_ORDER, V2_ORDER, V2_PAGE1_ROWS
from model import *

ROOT = Path(__file__).resolve().parents[1]
VD = ROOT / "vendors"
PLANT_NAME = {"BLR": "Bengaluru (Hoskote)", "PNQ": "Pune (Chakan)", "HYD": "Hyderabad (Patancheru)"}


def inr(n):
    """Indian digit grouping for integers."""
    s = str(int(n))
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts) + "," + tail


def short_desc(l, style="v1"):
    ply = l["ply"]
    size = l["dimensions_text"].replace(" ", "")
    cat = l["price_category"]
    prt = l["printing"]
    if style == "v1":
        base = {"PT": f"{ply} Ply Partition", "PAD": f"{ply} Ply Layer Pad"}.get(cat)
        if base:
            return f"{base} {size} – {l['item_description'].split('–')[-1].strip() if '–' in l['item_description'] else ''}".strip(" –")
        return f"{ply} Ply {l['style']} {size} ({l['flute']} Flute) – {prt}"
    return l["item_description"]


# ============================================================ VENDOR 1 – Excel
def write_v1():
    v = VENDOR_BY_ID["VND-1001"]
    t = TERMS["VND-1001"]
    out = VD / "vendor_1"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A1"] = "DECCAN CORRUPACK INDUSTRIES PVT. LTD."
    ws["A1"].font = Font(bold=True, size=15, color="7F3F00")
    ws["A2"] = "No. 112, 4th Phase, Peenya Industrial Area, Bengaluru 560058  |  GSTIN 29AADCD8812K1Z4  |  Works: Hoskote & Hosur"
    ws["A3"] = "COMMERCIAL OFFER – CORRUGATED BOXES"
    ws["A3"].font = Font(bold=True, size=12)
    ws["A4"] = f"Our Ref: {t['quote_ref']}"
    ws["F4"] = "Date: 08-Oct-2026"
    ws["A5"] = f"To: {BUYER['name']} – Attn. Ms. {BUYER['contact']['name']}, {BUYER['contact']['title']}"
    ws["A6"] = f"Your Enquiry: {EVENT['rfx_id']} dated 21-Sep-2026"
    ws["A7"] = "We thank you for the enquiry and are pleased to quote as under. Items grouped by delivery location."
    heads = ["S.No", "Your Item Code", "Description", "Size L x W x H (mm)", "Ply / Flute", "Board Construction (GSM/BF)",
             "Delivery To", "Annual Qty (Nos)", "Basic Rate (Rs/Nos)", "Freight (Rs/Nos)", "MOQ (Nos)", "Lead Time (Days)", "Remarks"]
    fill = PatternFill("solid", fgColor="F4B183")
    thin = Side(style="thin", color="999999")
    hr = V1_FIRST_ROW - 1
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    remarks = {
        "CP-012": "Aqueous WR coating, Cobb(30) <= 40",
        "CP-015": "Artwork approval required before plates",
        "CP-017": "Tear-off perforation as per SCPL drawing",
        "CP-019": "Hand holes die-cut; std lead time",
        "CP-020": "Rate per set",
        "CP-021": "Rate per set",
        "CP-022": "Rate per set",
        "CP-029": "Both liners WR coated",
        "CP-030": "Pallet-fit, 4 per pallet",
    }
    for n, l in enumerate(V1_ROW_ORDER, 1):
        r = hr + n
        qv = QUOTES["VND-1001"][l["line_id"]]
        rate = qv["price"]
        rate_val = f"{rate:.2f} (incl. tear strip)" if l["line_id"] == "CP-018" else rate
        row = [n, l["line_id"], short_desc(l), l["dimensions_text"].replace(" x ", "x"), f"{l['ply']} / {l['flute']}",
               "/".join(f"{g}{f or ''}" + (f"(BF{bf})" if bf else "(duplex)") for g, bf, f in CONSTRUCTIONS[l["construction_code"]]["papers"]), PLANT_NAME[l["delivery_plant"]],
               l["annual_qty"], rate_val, qv["freight_inr"], qv["moq"], qv["lead_time_days"], remarks.get(l["line_id"], "")]
        for i, val in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=val)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if i in (9, 10) and isinstance(val, float):
                c.number_format = "0.00"
            if i in (8, 11):
                c.number_format = "#,##0"
        if l["line_id"] == "CP-015":
            ws.cell(row=r, column=9).comment = Comment("Litho lamination done through our partner Prism Offset; cost included in rate.", "Ramesh G")
        if l["line_id"] == "CP-029":
            ws.cell(row=r, column=10).comment = Comment("Freight worked on full-truck lots for 7 ply export boxes.", "Ramesh G")
    last = hr + len(V1_ROW_ORDER)
    terms = [
        "TERMS & CONDITIONS",
        "1. Basic rates are ex-works Hoskote. Freight up to your respective plant shown separately per box in column J.",
        "2. GST extra as applicable (HSN 4819).",
        "3. Payment: 60 days from date of GRN.",
        "4. Validity: 180 days from date of this offer.",
        "5. Delivery: 12 days from PO for regular items; 18 days for die-cut items.",
        "6. Printing plates and cutting dies at our cost – no tooling charges.",
        "7. Annual rebate: 1% on total billing above Rs. 2.00 Crore (excl. GST) in the contract year, by credit note in April.",
        "8. Tolerance on quantity +/- 5% per lot.",
        "",
        "For Deccan Corrupack Industries Pvt. Ltd.",
        "Ramesh Gowda – Head, Key Accounts  |  ramesh.gowda@deccancorrupack.example  |  +91 98450 31277",
    ]
    for i, tline in enumerate(terms):
        c = ws.cell(row=last + 2 + i, column=1, value=tline)
        if i == 0:
            c.font = Font(bold=True)
    for i, wdt in enumerate([5, 11, 46, 15, 8, 42, 22, 12, 14, 12, 9, 10, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.row_dimensions[hr].height = 32
    ws.freeze_panes = ws.cell(row=V1_FIRST_ROW, column=3)

    tq = wb.create_sheet("Tech Questionnaire")
    tq.append(["Q.No", "Area", "Response", "Supporting document"])
    for c in tq[1]:
        c.font = Font(bold=True)
        c.fill = fill
    docs = {"Q01": "ISO cert copy attached", "Q02": "FSC cert attached", "Q05": "Sample COA attached", "Q10": "BCP summary attached", "Q13": "ISO 14001 cert attached"}
    for qq in QUESTIONS:
        tq.append([qq["qid"], qq["topic"], QA["VND-1001"][qq["qid"]][1], docs.get(qq["qid"], "")])
    for i, wdt in enumerate([6, 26, 100, 24], 1):
        tq.column_dimensions[get_column_letter(i)].width = wdt
    ci = wb.create_sheet("Company Info")
    for row in [["Company", v["vendor_name"]], ["Plants", ", ".join(v["plant_locations"])], ["Installed capacity", "36,000 MT per annum"],
                ["Certifications", ", ".join(v["certifications"])], ["Key contact", "Ramesh Gowda, ramesh.gowda@deccancorrupack.example"],
                ["Bank", "Details on request"], ["MSME", "Not applicable"]]:
        ci.append(row)
    ci.column_dimensions["A"].width = 20
    ci.column_dimensions["B"].width = 80
    wb.save(out / "Deccan_Corrupack_Quotation_SCPL_RFQ014.xlsx")


# ============================================================ VENDOR 2 – PDF
def write_v2():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = VD / "vendor_2"
    out.mkdir(parents=True, exist_ok=True)
    path = out / "Shree_Ganesh_Packwell_Quotation_SGP-Q-2026-1187.pdf"
    t = TERMS["VND-1002"]

    def letterhead(c, doc):
        c.saveState()
        W, H = A4
        c.setFillColor(colors.HexColor("#8B1A1A"))
        c.rect(0, H - 30 * mm, W, 30 * mm, stroke=0, fill=1)
        c.setFillColor(colors.HexColor("#F2C14E"))
        c.circle(22 * mm, H - 15 * mm, 9 * mm, stroke=0, fill=1)
        c.setFillColor(colors.HexColor("#8B1A1A"))
        c.setFont("Helvetica-Bold", 13)
        c.drawCentredString(22 * mm, H - 17 * mm, "SGP")
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 17)
        c.drawString(36 * mm, H - 14 * mm, "SHREE GANESH PACKWELL PVT. LTD.")
        c.setFont("Helvetica", 8)
        c.drawString(36 * mm, H - 19.5 * mm, "Manufacturers of Corrugated Boxes, Partitions & Pads  |  ISO 9001:2015  |  FSC-CoC")
        c.drawString(36 * mm, H - 24 * mm, "Plot A-27, Chakan MIDC Phase I, Taluka Khed, Pune 410501  |  GSTIN 27AAHCS5530P1ZB  |  sales@ganeshpackwell.example")
        c.setFillColor(colors.grey)
        c.setFont("Helvetica", 7)
        c.drawString(15 * mm, 10 * mm, "Regd. Office: 14, Shivaji Nagar, Pune 411005  |  CIN U21012PN2009PTC134221")
        c.drawRightString(W - 15 * mm, 10 * mm, f"Page {doc.page}")
        c.restoreState()

    st = ParagraphStyle("b", fontName="Helvetica", fontSize=9, leading=12)
    stb = ParagraphStyle("bb", parent=st, fontName="Helvetica-Bold")
    sth = ParagraphStyle("h", parent=st, fontName="Helvetica-Bold", fontSize=12, leading=15, spaceAfter=4)
    small = ParagraphStyle("s", parent=st, fontSize=6.5, leading=8, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("c", parent=st, fontSize=8, leading=10)

    def desc2(l):
        cat = l["price_category"]
        sz = l["dimensions_text"].replace(" ", "")
        if cat in ("PT",):
            return f"{l['ply']} Ply Partition {l['item_description'].split('–')[1].split(',')[0].strip()}"
        if cat == "PAD":
            return f"{l['ply']} Ply Pad {sz}"
        prt = "Plain" if l["printing"] == "Plain" else l["printing"].split()[0].replace("C", " Col") + " Printed"
        extra = " WR Coated" if l["coating"] and l["line_id"] == "CP-012" else ""
        sty = {"RSC": "RSC", "HSC": "HSC", "RSC export": "Export RSC"}.get(l["style"], l["style"].title())
        return f"{l['ply']} Ply {sty} {prt}{extra}"

    def rows(lines):
        data = [["Sl.", "Item Code", "Description", "Size (mm)", "Ply", "Rate (Rs.)", "Per", "Del. Plant"]]
        for i, l in lines:
            qv = QUOTES["VND-1002"][l["line_id"]]
            per = "100 Nos" if qv["pieces_per_unit"] == 100 else ("Set" if l["uom"] == "set" else "Nos")
            data.append([str(i), l["line_id"], Paragraph(desc2(l), cell), l["dimensions_text"].replace(" ", ""), str(l["ply"]),
                         f"{qv['price']:,.2f}", per, l["delivery_plant"]])
        return data

    def table(data):
        tb = Table(data, colWidths=[9 * mm, 18 * mm, 62 * mm, 26 * mm, 9 * mm, 20 * mm, 16 * mm, 18 * mm], repeatRows=1)
        tb.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3E3C3")), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#AAAAAA")),
            ("ALIGN", (5, 1), (5, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FBF7EF")]),
        ]))
        return tb

    indexed = list(enumerate(V2_ORDER, 1))
    story = [Spacer(1, 22 * mm),
             Table([[Paragraph("<b>QUOTATION</b>", sth), Paragraph(f"Ref: <b>{t['quote_ref']}</b><br/>Date: 10/10/2026", st)]], colWidths=[110 * mm, 68 * mm]),
             Paragraph(f"To,<br/>The Purchase Department – Packaging<br/>{BUYER['name']}<br/>Baner, Pune 411045<br/>Kind Attn: Ms. {BUYER['contact']['name']}", st),
             Spacer(1, 3 * mm),
             Paragraph(f"<b>Sub:</b> Quotation against your RFQ No. {EVENT['rfx_id']} for corrugated packaging", st),
             Spacer(1, 2 * mm),
             Paragraph("Dear Madam, with reference to the above enquiry we are pleased to submit our most competitive rates for the following items. "
                       "Rates are basic and subject to the terms printed at the end of this quotation*.", st),
             Spacer(1, 3 * mm),
             table(rows(indexed[:V2_PAGE1_ROWS])),
             Spacer(1, 2 * mm),
             Paragraph("<i>contd. on next page</i>", small),
             PageBreak(), Spacer(1, 22 * mm)]
    dup = LINE_BY_ID[V2_DUPLICATE_LINE]
    data2 = rows(indexed[V2_PAGE1_ROWS:])
    data2.append([str(len(V2_ORDER) + 1), dup["line_id"], Paragraph("5 Ply Printed Box (3 clr)", cell), dup["dimensions_text"].replace(" ", ""), "5",
                  f"{QUOTES['VND-1002'][V2_DUPLICATE_LINE]['duplicate_price']:,.2f}", "Nos", "PNQ"])
    story += [table(data2), Spacer(1, 4 * mm),
              Paragraph("<b>Payment:</b> 45 days from date of invoice.", st),
              Paragraph("<b>Delivery:</b> 15 days from receipt of PO / artwork approval.", st),
              Paragraph("<b>GST:</b> @ 5% extra.", st),
              Spacer(1, 3 * mm),
              Paragraph("We assure you of our best quality and service at all times and look forward to your valued order.", st),
              Spacer(1, 2 * mm),
              Paragraph("For <b>Shree Ganesh Packwell Pvt. Ltd.</b><br/><br/>Sachin Kulkarni<br/>Director – Sales  |  +91 98220 64519  |  sachin.k@ganeshpackwell.example", st),
              Spacer(1, 6 * mm),
              Paragraph("* Terms & notes:", small)]
    notes = [
        "1. Rates are FOR our Chakan works / your Chakan plant. For supplies to Bengaluru &amp; Hyderabad plants freight extra at actuals.",
        "2. Special discount of 2% on basic value applicable where SCPL annual order value exceeds Rs. 1.50 Crore (excl. GST).",
        "3. GST @ 5% extra as applicable on date of supply.",
        "4. Cutting dies and printing plates will be charged at actuals, one time.",
        "5. Validity: 60 days from date of quotation.",
        "6. MOQ 3,000 Nos per size per lot; partitions 5,000 sets.",
        "7. Tolerance +/- 10% on quantity; +/- 3 mm on size.",
        "8. CP-015 (offset litho) and 7-ply export items CP-029 / CP-030 not quoted – litho not in our range; BF 26 kraft not currently available.",
        "9. Subject to Pune jurisdiction.",
    ]
    story += [Paragraph(n, small) for n in notes]
    story += [PageBreak(), Spacer(1, 22 * mm), Paragraph("Annexure – Technical Compliance Statement", sth),
              Paragraph(f"Against SCPL RFQ {EVENT['rfx_id']} technical questionnaire.", st), Spacer(1, 3 * mm)]
    tq = [["Q.", "Area", "Our Response"]]
    for qq in QUESTIONS:
        tq.append([qq["qid"], qq["topic"], Paragraph(QA["VND-1002"][qq["qid"]][1], cell)])
    tqt = Table(tq, colWidths=[12 * mm, 38 * mm, 128 * mm])
    tqt.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
                             ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#AAAAAA")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3E3C3")),
                             ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [tqt]
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=12 * mm, bottomMargin=18 * mm,
                            title="Quotation SGP/Q/2026/1187", author="Shree Ganesh Packwell Pvt. Ltd.")
    doc.build(story, onFirstPage=letterhead, onLaterPages=letterhead)


# ============================================================ VENDOR 3 – DOCX
def write_v3():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    out = VD / "vendor_3"
    out.mkdir(parents=True, exist_ok=True)
    d = Document()
    stl = d.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)
    h = d.add_paragraph()
    r = h.add_run("INDUS FIBREBOARD SOLUTIONS LIMITED")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x0B, 0x4F, 0x6C)
    d.add_paragraph("Corporate Office: Survey No. 64, Gachibowli, Hyderabad 500032  |  CIN L21015TG2004PLC043781  |  GSTIN 36AACCI7123F1ZM\n"
                    "Plants: Patancheru (TS) · Ranjangaon (MH) · Hosur (TN)")
    d.add_paragraph(f"Ref: {TERMS['VND-1003']['quote_ref']}\t\t\t\t\t\tDate: 9 October 2026")
    d.add_paragraph(f"To\nMs. {BUYER['contact']['name']}\n{BUYER['contact']['title']}\n{BUYER['name']}\nBaner, Pune 411045")
    p = d.add_paragraph()
    p.add_run(f"Subject: Techno-commercial offer against RFQ {EVENT['rfx_id']} – Corrugated Packaging").bold = True
    d.add_paragraph(
        "Dear Ms. Raghunathan,\n\nThank you for inviting Indus Fibreboard Solutions to participate in your annual corrugated packaging tender. "
        "We have reviewed the specifications for all thirty line items and are pleased to submit our offer below. All three of our plants are "
        "located close to your manufacturing sites, which allows us to supply every location on a FOR basis with freight included in the rates.")

    d.add_heading("1. Price schedule – domestic supply", level=2)
    d.add_paragraph("Table 1: Rates in INR per piece (partitions per set), exclusive of GST, FOR your plant.")
    dom = [l for l in LINES if l["line_id"] not in COVERAGE_MISSING["VND-1003"] and l["line_id"] not in V3_USD_LINES]
    tb = d.add_table(rows=1, cols=7)
    tb.style = "Light Grid Accent 1"
    for i, hd in enumerate(["Item code", "Description", "Size (mm)", "Board", "Rate (INR)", "MOQ (pcs)", "Supplying plant"]):
        tb.rows[0].cells[i].text = hd
    sup = {"BLR": "Hosur", "PNQ": "Ranjangaon", "HYD": "Patancheru"}
    for l in dom:
        qv = QUOTES["VND-1003"][l["line_id"]]
        size = l["dimensions_text"]
        desc = l["item_description"]
        if l["line_id"] == V3_MISMATCH_LINE:
            size = "450 x 350 x 300"
            desc = "5 Ply Heavy Duty RSC (BF 22)"
        moq = "Refer note" if l["line_id"] == "CP-015" else f"{qv['moq']:,}"
        c = tb.add_row().cells
        for i, val in enumerate([l["line_id"], desc, size, f"{l['ply']}-ply {l['flute']}", f"{qv['price']:.2f}", moq, sup[l["delivery_plant"]]]):
            c[i].text = val
    d.add_paragraph()
    d.add_paragraph(
        "Please note that we are unable to offer CP-021 (24-cell partition) in this round, as our partition assembly line at Ranjangaon is fully "
        "booked until the second quarter of next year. We would be happy to quote for this item separately at that time.")

    d.add_heading("2. Export-grade cartons", level=2)
    u = {lid: QUOTES["VND-1003"][lid]["price"] for lid in V3_USD_LINES}
    d.add_paragraph(
        "Export-grade cartons are manufactured and invoiced by our Export Packaging Division, which bills in US dollars. "
        f"For the 5-ply export carton CP-028 (BF 24, two-colour with shipping marks) our price is USD {u['CP-028']:.2f} per piece. "
        f"The 7-ply moisture-resistant export carton CP-029 is offered at USD {u['CP-029']:.2f} per piece, and the heavy-duty pallet-fit "
        f"carton CP-030 at USD {u['CP-030']:.2f} per piece. These prices are delivered to the respective SCPL plants. Invoices will be raised in USD and "
        "converted at the RBI reference rate on the date of invoice (approximately INR 88.8 to the dollar today, for your reference only). "
        f"Minimum order quantity for export cartons is {QUOTES['VND-1003']['CP-028']['moq']:,} pieces per size.")

    d.add_heading("3. Commercial terms", level=2)
    t3 = d.add_table(rows=0, cols=2)
    t3.style = "Light List Accent 1"
    for k, val in [("Price basis", "FOR SCPL plants – freight included"), ("Currency", "INR (export cartons in USD – see Section 2)"),
                   ("Payment", "60 days from date of invoice"), ("Price validity", "90 days from date of offer"),
                   ("Taxes", "Prices exclusive of GST; GST as applicable"), ("Lead time", "14 days from PO"),
                   ("Tooling", "One-time INR 1,85,000 towards printing plates and cutting dies for printed and die-cut items"),
                   ("Quality documents", "COA with every lot; quarterly third-party test report")]:
        c = t3.add_row().cells
        c[0].text, c[1].text = k, val
    d.add_paragraph()
    d.add_paragraph(
        "Lead times: our standard lead time of 14 days applies to repeat orders. For the litho-laminated gift carton CP-015 and the "
        "shelf-ready display tray CP-017, the first supply will require 28 days from PO to allow for plate making, cutting-die "
        "manufacture and, in the case of CP-015, litho lamination proofing. Subsequent orders on these two items will follow the 14-day lead time.")
    d.add_paragraph(
        "Minimum order: for CP-015 we require a minimum production run of 25,000 pieces per artwork, owing to the make-ready involved on "
        "our litho-lamination line. MOQs for all other items are as shown in Table 1.")

    d.add_heading("4. Technical questionnaire responses", level=2)
    tq = d.add_table(rows=1, cols=3)
    tq.style = "Light Grid Accent 1"
    for i, hd in enumerate(["Q.", "Topic", "Response"]):
        tq.rows[0].cells[i].text = hd
    for qq in QUESTIONS:
        c = tq.add_row().cells
        c[0].text, c[1].text, c[2].text = qq["qid"], qq["topic"], QA["VND-1003"][qq["qid"]][1]
    d.add_paragraph()
    d.add_paragraph(
        "We believe our three-plant footprint, BRCGS-certified hygiene standards and in-house litho capability make us a strong long-term partner "
        "for Sahyadri. Given the expected revision in kraft paper prices from January, the above rates are valid until 31 December 2026.")
    d.add_paragraph("Warm regards,\n\nAnanya Reddy\nBusiness Development Manager – FMCG\nIndus Fibreboard Solutions Ltd.\n"
                    "ananya.reddy@indusfibreboard.example  |  +91 99490 22814")
    for para in d.paragraphs:
        para.paragraph_format.space_after = Pt(6)
    d.save(out / "Indus_Fibreboard_Offer_SCPL_RFQ014.docx")


# ============================================================ VENDOR 4 – printable rate card (HTML)
ROW_H = 30
V4_ROW_ORDER = [l for l in LINES if l["line_id"] not in COVERAGE_MISSING["VND-1004"]]


def write_v4_html():
    out = VD / "vendor_4"
    out.mkdir(parents=True, exist_ok=True)
    rows = []
    for i, l in enumerate(V4_ROW_ORDER, 1):
        lid = l["line_id"]
        qv = QUOTES["VND-1004"][lid]
        star = "*" if qv["pieces_per_unit"] == 100 else ""
        desc = l["item_description"].replace("Shipper ", "").replace(" – ", " - ")
        if len(desc) > 52:
            desc = desc[:50].rsplit(" ", 1)[0] + "…"
        rate = f"{qv['price']:.0f}{star}" if star else f"{qv['price']:.2f}"
        note = ""
        if lid == V4_NONCOMPLIANT_LINE:
            desc = "5 Ply RSC 2 Col (without WR coating)"
        if lid in ("CP-011", "CP-015"):
            note = '<span class="assoc">#</span>'
        cls = ' class="smudge"' if lid == V4_BLURRED_LINE else ""
        if lid == V4_HANDWRITTEN_LINE:
            rate_html = (f'<span class="struck">{qv["printed_struck_value"]:.2f}</span>'
                         f'<span class="hand">{qv["price"]:.2f}</span>')
        else:
            rate_html = rate
        rows.append(f'<tr data-line="{lid}"{cls}><td class="c">{i}</td><td class="c code">{lid}</td><td>{desc}{note}</td>'
                    f'<td class="c">{l["dimensions_text"].replace(" ", "")}</td><td class="c">{l["ply"]}</td><td class="r rate">{rate_html}</td></tr>')
    html = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Sri Vaigai Corrugators – Rate Card Oct 2026</title>
<style>
  @page {{ size: A4; margin: 10mm; }}
  body {{ margin:0; background:#fff; font-family: "Courier New", Courier, monospace; color:#1b1b1b; }}
  .sheet {{ width: 1000px; margin: 0 auto; padding: 28px 36px; background:#fbf9f2; }}
  .hdr {{ text-align:center; border-bottom: 3px double #333; padding-bottom: 8px; }}
  .hdr h1 {{ font-family: Georgia, "Times New Roman", serif; letter-spacing: 3px; margin: 0; font-size: 34px; color:#14365e; }}
  .hdr .sub {{ font-size: 14px; margin-top: 4px; }}
  .title {{ display:flex; justify-content:space-between; margin: 12px 0 8px; font-size: 15px; font-weight:bold; }}
  table {{ width:100%; border-collapse: collapse; font-size: 15px; }}
  th {{ background:#dfe6ee; border:1px solid #444; height: 36px; font-size: 14px; }}
  td {{ border:1px solid #666; height: {ROW_H - 2}px; padding: 0 8px; white-space: nowrap; overflow:hidden; }}
  td.c {{ text-align:center; }} td.r {{ text-align:right; font-weight:bold; }}
  td.code {{ font-weight:bold; }}
  .struck {{ text-decoration: line-through; text-decoration-thickness: 2px; color:#555; margin-right: 10px; }}
  .hand {{ font-family: "Bradley Hand", "Segoe Script", cursive; color:#1a3a9c; font-size: 20px; font-weight: bold; transform: rotate(-4deg); display:inline-block; }}
  .assoc {{ color:#a00; font-weight:bold; margin-left: 4px; }}
  .foot {{ margin-top: 10px; font-size: 13.5px; line-height: 1.45; }}
  .boxes {{ display:flex; gap: 18px; margin-top: 10px; }}
  .box {{ flex:1; border: 1.5px solid #333; padding: 8px 12px; font-size: 13.5px; line-height: 1.45; }}
  .box b.t {{ display:block; text-decoration: underline; margin-bottom: 4px; }}
  .sign {{ margin-top: 14px; display:flex; justify-content: space-between; font-size: 14px; }}
  .stamp {{ border: 2px solid #7a2a8a; color:#7a2a8a; border-radius: 50%; width: 120px; height: 120px; display:flex; align-items:center; justify-content:center;
            text-align:center; font-size: 11px; transform: rotate(-12deg); opacity: .75; }}
</style></head>
<body><div class="sheet">
  <div class="hdr"><h1>SRI VAIGAI CORRUGATORS</h1>
    <div class="sub">Corrugated Boxes · Partitions · Pads &nbsp;|&nbsp; No. 7, SIPCOT Phase-I, Hosur – 635126 (T.N.)</div>
    <div class="sub">GSTIN: 33ABMPM6620Q1Z7 &nbsp;|&nbsp; Udyam: UDYAM-TN-31-0045821 &nbsp;|&nbsp; Ph: 94431 50762</div></div>
  <div class="title"><span>RATE CARD – SAHYADRI CONSUMER PRODUCTS (RFQ 014)</span><span>Date: 07.10.2026</span></div>
  <table>
    <thead><tr><th style="width:52px">S.No</th><th style="width:96px">Your Code</th><th>Description</th><th style="width:150px">Size (mm)</th><th style="width:48px">Ply</th><th style="width:150px">Rate / Box (Rs.)</th></tr></thead>
    <tbody>
{chr(10).join(rows)}
    </tbody>
  </table>
  <div class="foot">* Rate per 100 Nos (partitions &amp; pads). &nbsp; # Printing through associate printer.<br>
  7 Ply export cartons (CP-029, CP-030) – not offered.</div>
  <div class="boxes">
    <div class="box"><b class="t">TERMS</b>
      Ex-factory Hosur. Freight extra.<br>GST extra.<br>Payment: 15 days from invoice.<br>Validity: 30 days.<br>
      Min. order 1,000 box per size.<br>Delivery: 7 days.<br>Printing block charges extra.</div>
    <div class="box"><b class="t">COMPANY PROFILE</b>
      Capacity: 6,500 MT / year.<br>ISO 9001: Applied – audit Dec 2026.<br>Testing: Bursting strength tester &amp; GSM balance. Test report on request.<br>
      Lot no. &amp; date printed on every box.<br>100% recycled kraft available.<br>Complaint: replacement within 7 days.<br>
      Flexo up to 2 colours in-house.</div>
  </div>
  <div class="sign"><div>For SRI VAIGAI CORRUGATORS<br><br><br>K. Murugan, Proprietor</div>
    <div class="stamp">SRI VAIGAI<br>CORRUGATORS<br>HOSUR</div></div>
</div></body></html>
"""
    (out / "Sri_Vaigai_Rate_Card_Oct2026.html").write_text(html, encoding="utf-8")


# ============================================================ VENDOR 5 – email
def write_v5():
    out = VD / "vendor_5"
    out.mkdir(parents=True, exist_ok=True)
    a = {qq["qid"]: QA["VND-1005"][qq["qid"]][1] for qq in QUESTIONS}
    txt = f"""From: P. Srinivas Rao <srinivas@godavaripack.example>
To: Meera Raghunathan <meera.raghunathan@sahyadricp.example>
Cc: purchase.hyd@sahyadricp.example
Date: Mon, 12 Oct 2026 18:47:12 +0530
Subject: RE: RFQ SCPL/PKG/RFQ/2026-27/014 - Corrugated boxes - our rates

Dear Madam,

Greetings from Godavari Packaging.

Thank you for the enquiry. Sorry for the delay, we were waiting for the revised kraft rates from the mill. Please find our rates below. All rates are per kg of box weight, GST extra.

- 3-ply plain brown boxes: Rs 44/kg
- 5-ply plain: Rs 46/kg (RSC and the big HSC box). For the heavy duty BF22 jam box add Rs 3/kg.
- Printed ones Rs 5/kg extra over plain, up to 2 colours - so the 1 & 2 colour 3-ply boxes and the 2 colour 5-ply shampoo box (CP-010).
- Partitions and layer pads (CP-020 to CP-024): Rs 41/kg.

For the trays and inserts we are already supplying to your Hyderabad plant, remaining items same as last year's pricing.

The 3-colour box CP-011 we cannot do (2 colour max), and the water resistant one CP-012 also we don't have coating facility. Other die-cut items (gift box, mailer, hand-hole box), the sleeve, export and 7 ply items we are not quoting this time.

Terms:
Freight extra.
MOQ 5,000 for printed boxes. Below that small run charges will apply.
Payment 30 days.
Delivery 10 days from PO.
Rates hold till December end. Rates are based on current kraft price, if kraft goes up more than 5% we will have to pass on.

Your technical questionnaire - our replies:
1. {a['Q01']}
2. {a['Q02']}
3. {a['Q03']}
4. {a['Q04']}
5. {a['Q05']}
6. {a['Q06']}
7. Lead time {a['Q07']}
8. {a['Q08']}
9. {a['Q09']}
10. {a['Q10']}
11. {a['Q11']}
12. {a['Q12']}
13. {a['Q13']}
14. {a['Q14']}

We have been supplying your Hyderabad plant for 5 years without any major complaint and request you to consider us for the Hyderabad and other requirements.

Regards,
P. Srinivas Rao
Partner, Godavari Packaging Works
Plot 118, IDA Bollaram, Sangareddy Dist., Telangana 502325
Mob: 98490 77312

Sent from my iPhone

-----Original Message-----
From: Meera Raghunathan
Sent: 21 September 2026 10:05
Subject: RFQ SCPL/PKG/RFQ/2026-27/014 - Corrugated boxes

Dear Supplier, please find attached our RFQ for corrugated packaging (30 line items) for our Bengaluru, Pune and Hyderabad plants. Kindly submit your quotation along with the technical questionnaire by 12 October 2026.
"""
    (out / "Godavari_Packaging_email_response.txt").write_text(txt, encoding="utf-8")


if __name__ == "__main__":
    write_v1()
    write_v2()
    write_v3()
    write_v4_html()
    write_v5()
    print("vendor docs written")
